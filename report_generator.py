"""
report_generator.py — сохранение результатов в .docx и .xlsx.
"""

import os
from datetime import datetime
from docx import Document
from openpyxl import Workbook


class ReportGenerator:
    """Генератор отчётов по результатам расчётов."""

    def __init__(self, output_dir: str = "reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def save_docx(self, data: dict, filename: str = None) -> str:
        """Сохраняет словарь результатов в .docx. Возвращает путь к файлу."""
        if filename is None:
            filename = f"report_{datetime.now():%Y%m%d_%H%M%S}.docx"
        path = os.path.join(self.output_dir, filename)

        doc = Document()
        doc.add_heading("Отчёт по расчёту геометрической фигуры", level=1)

        for key, value in data.items():
            p = doc.add_paragraph()
            p.add_run(f"{key}: ").bold = True
            p.add_run(str(value))

        doc.save(path)
        return path

    def save_xlsx(self, data: dict, filename: str = None) -> str:
        """Сохраняет словарь результатов в .xlsx. Возвращает путь к файлу."""
        if filename is None:
            filename = f"report_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        path = os.path.join(self.output_dir, filename)

        wb = Workbook()
        ws = wb.active
        ws.title = "Результаты"

        ws.append(["Параметр", "Значение"])
        for key, value in data.items():
            ws.append([key, str(value)])

        wb.save(path)
        return path